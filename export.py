# -*- coding: utf-8 -*-
"""导出 Excel 估值模型：预测、折现法、敏感性全部为活公式。
约定：蓝色字体 = 输入；黑色字体 = 公式。"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from i18n import T, TL

BLUE = Font(name="等线", size=10, color="0000CC")
BLACK = Font(name="等线", size=10)
BOLD = Font(name="等线", size=10, bold=True)
TITLE = Font(name="等线", size=14, bold=True)
NOTE = Font(name="等线", size=9, italic=True, color="808080")
HEAD = PatternFill("solid", fgColor="E7E6E6")
KEY = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
N2, N0, N4, PCT = "#,##0.00", "#,##0", "0.0000", "0.00"


def _c(ws, r, c, v, font=BLACK, fmt=None, fill=None, al=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font, x.border = font, BOX
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    if al:
        x.alignment = Alignment(horizontal=al, vertical="center", wrap_text=True)
    return x


def _head(ws, r, labels, c0=1):
    for i, t in enumerate(TL(labels)):
        _c(ws, r, c0 + i, t, BOLD, fill=HEAD, al="center")


def _sheet(name: str) -> str:
    """Excel 表名不能含 / \\ ? * [ ] :，且不超过 31 个字符。"""
    t = T(name)
    for ch in '/\\?*[]:':
        t = t.replace(ch, "")
    return t[:31]


# 「预测与折现法」表的行号（敏感性表与校验脚本按名称引用）
ROW = {"g": 5, "rev": 6, "m": 7, "core": 8, "other": 9, "pbt": 10, "np": 11, "nopat": 12,
       "da_pct": 13, "da": 14, "cx_pct": 15, "cx": 16, "dnwc": 17, "fcff": 18, "frac": 19, "t": 20, "disc": 21, "pv": 22,
       "ke": 25, "kd": 26, "wd": 27, "wacc": 29, "gt": 30, "pvsum": 33, "tv": 34, "tvpv": 35, "ev": 36,
       "cash": 37, "fin": 38, "debt": 39, "eq": 40, "sh": 41, "ps": 42, "tvp": 43, "vs": 44,
       "m_ev0": 48, "m_ev1": 49, "m_pe0": 50, "m_pe1": 51, "m_tv": 52, "chk": 56}
DRIVER_ROWS = ("g", "m", "other", "da_pct", "cx_pct")     # 逐年假设所在行
MULT = '0.0"x"'
GOOD = PatternFill("solid", fgColor="D9EAD3")
WARN = PatternFill("solid", fgColor="FCE8B2")
BAD = PatternFill("solid", fgColor="F4CCCC")


def write_dcf(wb, R: dict, SF: str, SS: str, base_label: str, base_cells: dict,
              base_year: int | None = None, n: int = 5, yr_cells: dict | None = None,
              extra_checks: list | None = None) -> None:
    """写入「预测与折现法」「敏感性」两张表。
    R：假设表单元格的绝对引用；base_cells：基期列（B 列）{行名: (值或公式, 字体)}；
    yr_cells：逐年假设 {行名: [(值或公式, 字体) × n]}，缺省时按假设表的首末年线性插值（空白模板用）；
    extra_checks：额外的模型检查 [(名称, 返回「通过」等文字的公式, 说明)]。
    预测期之后一列为「终值年」：最后一个预测年按永续增长率外推一年，资本开支取终值年假设。"""
    from openpyxl.formatting.rule import CellIsRule
    F = f"'{SF}'!"
    X = ROW
    wf = wb.create_sheet(SF)
    wf.column_dimensions["A"].width = 34
    for col in range(2, 4 + n):
        wf.column_dimensions[get_column_letter(col)].width = 14
    wf["A1"] = T("盈利预测与折现法"); wf["A1"].font = TITLE
    wf["A2"] = T("单位：万元；蓝色为逐年假设，可直接修改"); wf["A2"].font = NOTE
    yl = [f"{base_year + i + 1}E" for i in range(n)] if base_year else [f"E{i + 1}" for i in range(n)]
    _head(wf, 4, ["项目", base_label] + yl + ["终值年"])
    lab = {"g": "营收增速（%）", "rev": "营业总收入", "m": "核心经营利润率（%）", "core": "核心经营利润",
           "other": "投资收益及其他", "pbt": "利润总额", "np": "归母净利润", "nopat": "NOPAT",
           "da_pct": "折旧摊销 / 营收（%）", "da": "折旧摊销", "cx_pct": "资本开支 / 营收（%）", "cx": "资本开支",
           "dnwc": "营运资金增加", "fcff": "自由现金流 FCFF", "frac": "计入比例", "t": "折现年数",
           "disc": "折现因子", "pv": "现值"}
    key_rows = ("rev", "np", "fcff")
    for k, t in lab.items():
        _c(wf, X[k], 1, T(t), BOLD if k in key_rows + ("pv",) else BLACK,
           fill=KEY if k in key_rows else None, al="left")
    for k, (v, font) in base_cells.items():
        _c(wf, X[k], 2, v, font, PCT if k in ("g", "m", "da_pct", "cx_pct") else N2)
    if yr_cells is None:        # 按首末年插值（公式，黑色）
        yr_cells = {
            "g": [(f"={R['g1']}+({R['g5']}-{R['g1']})*{i}/{n - 1}", BLACK) for i in range(n)],
            "m": [(f"={R['m1']}+({R['m5']}-{R['m1']})*{i}/{n - 1}", BLACK) for i in range(n)],
            "other": [(f"={R['o1']}*(1+{R['oc']}/100)^{i}", BLACK) for i in range(n)],
            "da_pct": [(f"={R['da']}", BLACK)] * n,
            "cx_pct": [(f"={R['cx']}+({R['cx_ty']}-{R['cx']})*{i}/{n - 1}", BLACK) for i in range(n)]}

    def body(col: int, P: str):
        """一列的利润与现金流公式（预测年与终值年共用）。"""
        C = get_column_letter(col)
        r = lambda k: f"{C}{X[k]}"
        _c(wf, X["core"], col, f"={r('rev')}*{r('m')}/100", fmt=N2)
        _c(wf, X["pbt"], col, f"={r('core')}+{r('other')}", fmt=N2)
        _c(wf, X["np"], col, f"={r('pbt')}*(1-{R['tax']}/100)*(1-{R['mino']}/100)", fmt=N2, fill=KEY)
        _c(wf, X["nopat"], col, f"=({r('core')}+{r('other')}*{R['of']}/100)*(1-{R['tax']}/100)", fmt=N2)
        _c(wf, X["da"], col, f"={r('rev')}*{r('da_pct')}/100", fmt=N2)
        _c(wf, X["cx"], col, f"={r('rev')}*{r('cx_pct')}/100", fmt=N2)
        _c(wf, X["dnwc"], col, f"=({r('rev')}-{P}{X['rev']})*{R['nwc']}/100", fmt=N2)
        _c(wf, X["fcff"], col, f"={r('nopat')}+{r('da')}-{r('cx')}-{r('dnwc')}", fmt=N2, fill=KEY)

    for i in range(n):
        C = get_column_letter(3 + i); P = get_column_letter(2 + i)
        r = lambda k: f"{C}{X[k]}"
        for k in DRIVER_ROWS:
            v, font = yr_cells[k][i]
            _c(wf, X[k], 3 + i, v, font, N2 if k == "other" else PCT)
        _c(wf, X["rev"], 3 + i, f"={P}{X['rev']}*(1+{r('g')}/100)", fmt=N2, fill=KEY)
        body(3 + i, P)
        # 估值基准日之后的部分：第 1 年计入 stub，之后每年 1；年中折现取各期中点
        _c(wf, X["frac"], 3 + i, f"={R['stub']}" if i == 0 else 1, fmt="0%")
        _c(wf, X["t"], 3 + i, (f"=IF({R['mid']}=1,{R['stub']}/2,{R['stub']})" if i == 0
                               else f"={R['stub']}+{i}-IF({R['mid']}=1,0.5,0)"), fmt="0.00")
        _c(wf, X["disc"], 3 + i, f"=1/(1+$B${X['wacc']}/100)^{r('t')}", fmt=N4)
        _c(wf, X["pv"], 3 + i, f"={r('fcff')}*{r('frac')}*{r('disc')}", fmt=N2)

    LC = get_column_letter(2 + n)  # 最后一个预测年所在列
    TC = get_column_letter(3 + n)  # 终值年所在列
    ty = lambda k: f"{TC}{X[k]}"
    _c(wf, X["g"], 3 + n, f"={R['gt']}", fmt=PCT)
    _c(wf, X["rev"], 3 + n, f"={LC}{X['rev']}*(1+{ty('g')}/100)", fmt=N2, fill=KEY)
    _c(wf, X["m"], 3 + n, f"={LC}{X['m']}", fmt=PCT)
    _c(wf, X["other"], 3 + n, f"={LC}{X['other']}*(1+{ty('g')}/100)", fmt=N2)
    _c(wf, X["da_pct"], 3 + n, f"={LC}{X['da_pct']}", fmt=PCT)
    _c(wf, X["cx_pct"], 3 + n, f"={R['cx_ty']}", fmt=PCT)
    body(3 + n, LC)
    for k in ("frac", "t", "disc", "pv"):
        _c(wf, X[k], 3 + n, None)

    B = lambda k: f"B{X[k]}"
    ebitda0 = f"({B('core')}+{B('rev')}*{B('da_pct')}/100)"
    ebitda1 = f"(C{X['core']}+C{X['da']})"
    ebitda_t = f"({ty('core')}+{ty('da')})"
    mev = f"({R['mcap']}-({B('cash')}+{B('fin')}-{B('debt')}))"
    blk = [("ke", "股权成本（%）", f"={R['rf']}+{R['beta']}*{R['erp']}+{R['sp']}", PCT),
           ("kd", "税后债务成本（%）", f"={R['kd']}*(1-{R['tax']}/100)", PCT),
           ("wd", "债务权重（%）", f"=IF({R['debt']}+{R['mcap']}>0,{R['debt']}/({R['debt']}+{R['mcap']})*100,0)", PCT),
           ("wacc", "折现率 WACC（%）", f"={B('ke')}*(1-{B('wd')}/100)+{B('kd')}*{B('wd')}/100", PCT),
           ("gt", "永续增长率（%）", f"={R['gt']}", PCT),
           ("pvsum", "预测期现值合计", f"=SUM(C{X['pv']}:{LC}{X['pv']})", N2),
           ("tv", "终值（终值年 FCFF ÷ (WACC − g)）", f"={ty('fcff')}/({B('wacc')}/100-{B('gt')}/100)", N2),
           ("tvpv", "终值现值", f"={B('tv')}*{LC}{X['disc']}", N2),
           ("ev", "企业价值", f"={B('pvsum')}+{B('tvpv')}", N2),
           ("cash", "加：货币资金", f"={R['cash']}*{R['add_cash']}", N2),
           ("fin", "加：交易性金融资产", f"={R['fin']}*{R['add_fin']}", N2),
           ("debt", "减：有息负债", f"={R['debt']}", N2),
           ("eq", "股权价值", f"={B('ev')}+{B('cash')}+{B('fin')}-{B('debt')}", N2),
           ("sh", "总股本（万股）", f"={R['shares']}", N2),
           ("ps", "每股价值（元）", f"={B('eq')}/{B('sh')}", N2),
           ("tvp", "终值占企业价值（%）", f"={B('tvpv')}/{B('ev')}*100", PCT),
           ("vs", "较现价", f"={B('ps')}/{R['price']}-1", "0.0%")]
    wf.cell(row=X["ke"] - 1, column=1, value=T("折现率")).font = TITLE
    wf.cell(row=X["pvsum"] - 1, column=1, value=T("估值")).font = TITLE
    for k, t, f, fm in blk:
        hi = k in ("wacc", "ps")
        _c(wf, X[k], 1, T(t), BOLD if k in ("wacc", "ev", "eq", "ps") else BLACK, al="left",
           fill=KEY if hi else None)
        _c(wf, X[k], 2, f, fmt=fm, fill=KEY if hi else None)

    # ── 隐含倍数交叉检验：折现法结果换算成倍数，与现价对应的倍数并列
    nm = '"n.m."'
    wf.cell(row=X["m_ev0"] - 2, column=1, value=T("隐含倍数交叉检验")).font = TITLE
    _head(wf, X["m_ev0"] - 1, ["倍数", "折现法隐含", "按现价"])
    y0, y1 = base_label, yl[0]
    mrows = [("m_ev0", f"EV/EBITDA（{y0}）", f"=IF(AND({B('ev')}>0,{ebitda0}>0),{B('ev')}/{ebitda0},{nm})", f"=IF(AND({mev}>0,{ebitda0}>0),{mev}/{ebitda0},{nm})"),
             ("m_ev1", f"EV/EBITDA（{y1}）", f"=IF(AND({B('ev')}>0,{ebitda1}>0),{B('ev')}/{ebitda1},{nm})", f"=IF(AND({mev}>0,{ebitda1}>0),{mev}/{ebitda1},{nm})"),
             ("m_pe0", f"{T('市盈率')}（{y0}）", f"=IF(AND({B('eq')}>0,{B('np')}>0),{B('eq')}/{B('np')},{nm})",
              f"=IF({B('np')}>0,{R['mcap']}/{B('np')},{nm})"),
             ("m_pe1", f"{T('市盈率')}（{y1}）", f"=IF(AND({B('eq')}>0,C{X['np']}>0),{B('eq')}/C{X['np']},{nm})",
              f"=IF(C{X['np']}>0,{R['mcap']}/C{X['np']},{nm})"),
             ("m_tv", T("终值隐含 EV/EBITDA（终值年）"), f"=IF(AND({B('tv')}>0,{ebitda_t}>0),{B('tv')}/{ebitda_t},{nm})", "—")]
    for k, t, f1, f2 in mrows:
        _c(wf, X[k], 1, t, al="left")
        _c(wf, X[k], 2, f1, fmt=MULT, al="right")
        _c(wf, X[k], 3, f2, fmt=MULT, al="right")
    _c(wf, X["m_tv"] + 1, 1, T("EBITDA = 核心经营利润 + 折旧摊销；按现价的企业价值 = 总市值 − 净现金"), NOTE, al="left")
    wf.cell(row=X["m_tv"] + 1, column=1).border = Border()

    # ── 模型检查
    ok, warn, bad = T("通过"), T("需关注"), T("未通过")
    q = lambda x: f'"{x}"'
    checks = [("折现率高于永续增长率 3 个百分点以上",
               f"=IF({B('wacc')}-{B('gt')}>=3,{q(ok)},IF({B('wacc')}>{B('gt')},{q(warn)},{q(bad)}))",
               "差距过小时终值被大幅放大"),
              ("终值年自由现金流为正", f"=IF({ty('fcff')}>0,{q(ok)},{q(bad)})", "为负时终值没有意义"),
              ("终值年资本开支不低于折旧摊销", f"=IF({ty('cx_pct')}>={ty('da_pct')}-0.0001,{q(ok)},{q(warn)})",
               "稳定状态下资本开支至少维持现有资产"),
              ("终值占企业价值不超过 85%", f"=IF({B('tvp')}<=85,{q(ok)},{q(warn)})", "过高说明估值主要依赖永续假设")]
    checks += list(extra_checks or [])
    c0 = X["chk"]
    wf.cell(row=c0 - 2, column=1, value=T("模型检查")).font = TITLE
    _head(wf, c0 - 1, ["检查项", "结果", "说明"])
    has = f"ISNUMBER({B('ps')})"          # 还没有估值结果（空白模板未填数据）时检查项留空
    for j, (t, f, note) in enumerate(checks):
        _c(wf, c0 + j, 1, T(t), al="left")
        _c(wf, c0 + j, 2, f'=IF({has},{f[1:]},"")', al="center")
        _c(wf, c0 + j, 3, T(note), NOTE, al="left")
    last = c0 + len(checks) - 1
    rng = f"B{c0}:B{last}"
    _c(wf, last + 1, 1, T("模型检查汇总"), BOLD, fill=KEY, al="left")
    _c(wf, last + 1, 2, f'=IF({has},IF(COUNTIF({rng},{q(ok)})={len(checks)},{q(T("全部通过"))},'
                        f'COUNTIF({rng},{q(ok)})&"/{len(checks)} "&{q(T("项通过"))}),"")', BOLD, fill=KEY, al="center")
    wf.merge_cells(start_row=c0 - 1, start_column=3, end_row=c0 - 1, end_column=6)
    for j in range(len(checks)):
        wf.merge_cells(start_row=c0 + j, start_column=3, end_row=c0 + j, end_column=6)
    wf.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[q(ok)], fill=GOOD))
    wf.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[q(warn)], fill=WARN))
    wf.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[q(bad)], fill=BAD))

    # ── 敏感性：每股价值、终值隐含 EV/EBITDA、隐含市盈率（同一网格，坐标轴随「假设」联动）
    wsn = wb.create_sheet(SS)
    wsn.column_dimensions["A"].width = 22
    for col in "BCDEF":
        wsn.column_dimensions[col].width = 12
    wsn["A1"] = T("敏感性分析"); wsn["A1"].font = TITLE
    wsn["A2"] = T("行：折现率（%）；列：永续增长率（%）。坐标轴随「假设」联动。"); wsn["A2"].font = NOTE
    dws, dgs = [-2, -1, 0, 1, 2], [-1.0, -0.5, 0, 0.5, 1.0]
    fr = X["fcff"]
    a_ = lambda k, c=LC: f"{F}${c}${X[k]}"
    sh_ = f"{F}$B${X['sh']}"

    def fty(g):       # 终值年 FCFF（随永续增长率 g 变化）
        return (f"(({a_('nopat')}+{a_('rev')}*({a_('da_pct')}-{a_('cx_pct', TC)})/100)*(1+{g}/100)"
                f"-{a_('rev')}*{g}/100*{R['nwc']}/100)")

    def grid(top: int, title: str, cell, fmt):
        wsn.cell(row=top - 1, column=1, value=title).font = BOLD
        _c(wsn, top, 1, "WACC \\ g", BOLD, fill=HEAD, al="center")
        for j, dg in enumerate(dgs):
            _c(wsn, top, 2 + j, f"={F}$B${X['gt']}+({dg})", BOLD, PCT, HEAD, "center")
        for i, dw in enumerate(dws):
            r_ = top + 1 + i
            _c(wsn, r_, 1, f"={F}$B${X['wacc']}+({dw})", BOLD, PCT, HEAD, "center")
            for j in range(len(dgs)):
                _c(wsn, r_, 2 + j, cell(f"$A{r_}", f"{get_column_letter(2 + j)}${top}", r_, j), fmt=fmt,
                   fill=KEY if (dw == 0 and dgs[j] == 0) else None, al="center")

    pv_ = lambda w: (f"SUMPRODUCT({F}$C${fr}:${LC}${fr},{F}$C${X['frac']}:${LC}${X['frac']},"
                     f"1/(1+{w}/100)^{F}$C${X['t']}:${LC}${X['t']})")
    grid(4, T("每股价值（元）"), lambda w, g, r_, j: (
        f"=IF({w}<={g},\"—\",({pv_(w)}+{fty(g)}/({w}/100-{g}/100)/(1+{w}/100)^{a_('t')}"
        f"+{F}$B${X['cash']}+{F}$B${X['fin']}-{F}$B${X['debt']})/{sh_})"), N2)
    ebt = lambda g: f"(({a_('core')}+{a_('da')})*(1+{g}/100))"
    grid(12, T("终值隐含 EV/EBITDA（倍）"), lambda w, g, r_, j: (
        f"=IF({w}<={g},\"—\",IF(OR({ebt(g)}<=0,{fty(g)}<=0),\"n.m.\",{fty(g)}/({w}/100-{g}/100)/{ebt(g)}))"), MULT)
    grid(20, T("隐含市盈率（{y}，倍）", y=yl[0]), lambda w, g, r_, j: (
        f"=IF(ISNUMBER({get_column_letter(2 + j)}{r_ - 16}),IF(AND({get_column_letter(2 + j)}{r_ - 16}>0,{F}$C${X['np']}>0),"
        f"{get_column_letter(2 + j)}{r_ - 16}*{sh_}/{F}$C${X['np']},\"n.m.\"),\"—\")"), MULT)


def build_workbook(cd: dict, A, res: dict) -> bytes:
    q, an, bs = cd["quote"], cd["annual"], cd["bs_latest"]
    base = int(an.index.max())
    wb = Workbook()
    wb.remove(wb.active)
    SA, SF, SS = [_sheet(x) for x in ("假设", "预测与折现法", "敏感性")]
    n = 5

    # ── 说明
    ws = wb.create_sheet(_sheet("说明"))
    ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 90
    ws["A1"] = T("{n}（{c}）估值模型", n=q["name"], c=q["code"]); ws["A1"].font = TITLE
    rows = [("数据时间", cd.get("asof", "")),
            ("单位", "金额：万元；每股：元；比率：%"),
            ("格式约定", "蓝色字体为输入，黑色字体为公式。改动「假设」表中任一蓝色单元格，全部结果自动重算。"),
            ("工作表", "假设 → 预测与折现法（含隐含倍数与模型检查）→ 敏感性；历史数据与蒙特卡洛结果为数值。"),
            ("模型结构", "营业总收入 × 核心经营利润率 = 核心经营利润；+ 投资收益及其他 = 利润总额；"
                     "× (1 − 有效税率) × (1 − 少数股东占比) = 归母净利润。"),
            ("折现时点", "估值基准日取最新报表日（与净现金同一时点），第 1 年只计入基准日之后的部分；默认年中折现。"),
            ("终值", "资本开支 / 营收从第 1 年线性过渡到稳定水平（第 5 年）；「终值年」= 第 5 年按永续增长率外推一年，"
                   "营运资金按永续增速计算；终值 = 终值年 FCFF ÷ (WACC − g)。"),
            ("交叉检验", "「预测与折现法」表下方：折现法隐含的 EV/EBITDA、市盈率与现价倍数并列，以及模型检查结果。"),
            ("数据来源", "东方财富（报表、行情、国债）、腾讯（日线、行情备用）、新浪（日线备用）公开接口，经 akshare 与直接请求获取。非官方授权数据。"),
            ("声明", "个人学习项目，结论基于用户假设，不构成投资建议。")]
    for i, (k, v) in enumerate(rows, start=3):
        _c(ws, i, 1, T(k), BOLD, fill=KEY); _c(ws, i, 2, T(v), al="left")

    # ── 假设（命名单元格位置）
    wa = wb.create_sheet(SA)
    wa.column_dimensions["A"].width = 30; wa.column_dimensions["B"].width = 14
    wa.column_dimensions["C"].width = 60
    wa["A1"] = T("假设与输入"); wa["A1"].font = TITLE
    _head(wa, 3, ["项目", "数值", "说明"])
    items = [
        ("base_rev", T("{y}A 营业总收入（万元）", y=base), float(an.loc[base, "营业总收入"]), "财报"),
        ("of", "其中计入经营现金流比例（%）", A.other_in_fcf, "0 表示折现法只对核心主业定价"),
        ("tax", "有效税率（%）", A.tax, ""),
        ("mino", "少数股东占比（%）", A.minority, ""),
        ("nwc", "营运资金增加 / 营收增量（%）", A.nwc_pct, "营收增速、利润率、投资收益、折旧摊销与资本开支为逐年假设，见「预测与折现法」表蓝色行"),
        ("rf", "无风险利率（%）", A.rf, "10年期国债"),
        ("beta", "贝塔", A.beta, "100周回归"),
        ("erp", "市场风险溢价（%）", A.erp, ""),
        ("sp", "规模溢价（%）", A.size_prem, ""),
        ("kd", "税前债务成本（%）", A.kd, ""),
        ("gt", "永续增长率（%）", A.g_term, ""),
        ("cx_ty", "第5年及终值年资本开支 / 营收（%）", A.capex_ty, "稳定状态的维持性资本开支，默认等于折旧摊销 / 营收；预测期从第 1 年线性过渡到此值"),
        ("mid", "年中折现（1=是，0=否）", int(A.mid_year), "现金流按各期中点折现"),
        ("stub", "第 1 年计入比例", A.stub, T("STUB_NOTE", d=bs["报告期"])),
        ("cash", "货币资金（万元）", bs["货币资金"], bs["报告期"]),
        ("fin", "交易性金融资产（万元）", bs["交易性金融资产"], bs["报告期"]),
        ("debt", "有息负债（万元）", bs["有息负债"], bs["报告期"]),
        ("add_cash", "货币资金计入净现金（1=是，0=否）", int(A.add_cash), ""),
        ("add_fin", "交易性金融资产计入净现金（1=是，0=否）", int(A.add_fin), ""),
        ("shares", "总股本（万股）", q["shares"] / 1e4, "行情接口"),
        ("mcap", "总市值（万元）", q["mcap_yi"] * 1e4, "行情接口"),
        ("price", "现价（元）", q["price"], "行情接口"),
    ]
    R = {}
    for i, (k, lab, v, note) in enumerate(items, start=4):
        _c(wa, i, 1, T(lab), al="left")
        _c(wa, i, 2, float(v), BLUE, N2 if abs(float(v)) >= 100 else "0.00")
        _c(wa, i, 3, T(note), NOTE, al="left")
        R[k] = f"'{SA}'!$B${i}"

    L = an.loc[base]
    drv = res["drv"]
    write_dcf(wb, R, SF, SS, f"{base}A", {
        "g": (float(L["营收增速%"]), BLUE), "rev": (f"={R['base_rev']}", BLACK),
        "m": (float(L["核心经营利润率%"]), BLUE), "core": (float(L["核心经营利润"]), BLUE),
        "other": (float(L["投资收益及其他"]), BLUE), "pbt": (float(L["利润总额"]), BLUE),
        "np": (float(L["归母净利润"]), BLUE), "da_pct": (float(L["折旧摊销/营收%"]), BLUE),
        "cx_pct": (float(L["资本开支/营收%"]), BLUE)},
        base_year=base, yr_cells={k: [(float(x), BLUE) for x in drv[src]] for k, src in
                                  (("g", "g"), ("m", "m"), ("other", "other"), ("da_pct", "da"), ("cx_pct", "capex"))})

    # ── 历史数据
    wh = wb.create_sheet(_sheet("历史数据"))
    wh.column_dimensions["A"].width = 24
    wh["A1"] = T("历史财务数据（万元 / %）"); wh["A1"].font = TITLE
    show = ["营业总收入", "营收增速%", "毛利率%", "营业总成本", "核心经营利润", "核心经营利润率%",
            "投资收益及其他", "利润总额", "有效税率%", "归母净利润", "扣非归母净利润", "非经常性损益",
            "研发费用", "经营现金流", "资本开支", "折旧摊销", "货币资金", "交易性金融资产",
            "有息负债", "资产总计", "归母权益"]
    _head(wh, 3, ["项目"] + [f"{y}A" for y in an.index])
    for i, k in enumerate(show, start=4):
        _c(wh, i, 1, T(k), al="left")
        for j, y in enumerate(an.index):
            v = an.loc[y, k]
            _c(wh, i, 2 + j, None if v != v else float(v), BLUE, N2)
    for j in range(len(an.index)):
        wh.column_dimensions[get_column_letter(2 + j)].width = 13

    # ── 蒙特卡洛
    wm = wb.create_sheet(_sheet("蒙特卡洛"))
    wm.column_dimensions["A"].width = 30
    for col in "BCDEF":
        wm.column_dimensions[col].width = 13
    wm["A1"] = T("蒙特卡洛模拟（数值）"); wm["A1"].font = TITLE
    e = res["emp"]
    info = [("参数来源", e.get("方法", "—")), ("观测区间", e.get("区间", "—")), ("观测数", e.get("观测数", 0)),
            ("营收增速标准差（pp）", e.get("增速标准差")), ("核心经营利润率标准差（pp）", e.get("利润率标准差")),
            ("相关系数", e.get("相关系数")), ("折现率σ（pp）", res.get("sig_w", 0.0)),
            ("永续增长率σ（pp）", res.get("sig_gt", 0.0)), ("第1年增速波动缩放", res["y1_scale"]),
            ("模拟次数", res["sims"])]
    for i, (k, v) in enumerate(info, start=3):
        _c(wm, i, 1, T(k), al="left"); _c(wm, i, 2, v, BLUE, N2 if isinstance(v, float) else None)
    _head(wm, 14, ["年份"] + [f"{base + i + 1}E" for i in range(n)])
    _c(wm, 15, 1, T("增速σ（pp）"), al="left"); _c(wm, 16, 1, T("利润率σ（pp）"), al="left")
    for i in range(n):
        _c(wm, 15, 2 + i, float(res["mc_sg"][i]), BLUE, N2)
        _c(wm, 16, 2 + i, float(res["mc_sm"][i]), BLUE, N2)
    _head(wm, 18, ["指标", "P5", "P25", "P50", "P75", "P95"])
    for r_, (k, d) in enumerate([(T("{y}年归母净利润（万元）", y=base + 1), res["mc_np"]),
                                  (T("折现法每股价值（元）"), res["mc_ps"])], start=19):
        _c(wm, r_, 1, k, al="left")
        for j, p in enumerate(["P5", "P25", "P50", "P75", "P95"]):
            _c(wm, r_, 2 + j, float(d[p]), BLUE, N2)
    wm["A22"] = T("蒙特卡洛依赖随机抽样，无法用单元格公式表达，此处为导出时的计算结果。")
    wm["A22"].font = NOTE

    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
