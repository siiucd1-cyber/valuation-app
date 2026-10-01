# -*- coding: utf-8 -*-
"""折现法 Excel 模板：填入 5 年历史报表科目 → 自动算出历史指标与参考假设 → 填写假设 → 折现法与敏感性自动计算。
不依赖网页，可单独发给别人使用。传入 example 时用一家公司的真实数据预填，便于核对公式。"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import i18n
from export import BLACK, BLUE, BOLD, KEY, N2, NOTE, PCT, TITLE, _c, _head, write_dcf

HS, AS, FS, SS = "历史数据", "假设", "预测与折现法", "敏感性"
YEARS = 5
COLS = [get_column_letter(2 + i) for i in range(YEARS)]      # B..F
LAST = COLS[-1]

# 历史数据输入行：(行号, 名称, 报表与科目, 预填字段)
INPUTS = [
    (5, "年份", "", None),
    (6, "营业总收入", "利润表：营业总收入", "营业总收入"),
    (7, "营业总成本", "利润表：营业总成本（含营业成本、税金及附加、销售/管理/研发/财务费用）", "营业总成本"),
    (8, "利润总额", "利润表：利润总额", "利润总额"),
    (9, "所得税费用", "利润表：所得税费用", "所得税"),
    (10, "净利润", "利润表：净利润", "净利润"),
    (11, "归母净利润", "利润表：归属于母公司股东的净利润", "归母净利润"),
    (12, "扣非归母净利润", "附注：扣除非经常性损益后归属于母公司股东的净利润", "扣非归母净利润"),
    (13, "折旧与摊销", "现金流量表补充资料：固定资产折旧 + 无形资产摊销 + 长期待摊费用摊销 + 使用权资产折旧", "折旧摊销"),
    (14, "资本开支", "现金流量表：购建固定资产、无形资产和其他长期资产支付的现金", "资本开支"),
    (15, "应收账款", "资产负债表：应收账款", "应收账款"),
    (16, "存货", "资产负债表：存货", "存货"),
    (17, "应付账款", "资产负债表：应付账款", "应付账款"),
    (18, "货币资金", "资产负债表：货币资金", "货币资金"),
    (19, "交易性金融资产", "资产负债表：交易性金融资产", "交易性金融资产"),
    (20, "短期借款", "资产负债表：短期借款", "短期借款"),
    (21, "一年内到期的非流动负债", "资产负债表：一年内到期的非流动负债", "一年内到期非流动负债"),
    (22, "长期借款", "资产负债表：长期借款", "长期借款"),
    (23, "应付债券", "资产负债表：应付债券", "应付债券"),
]


def _ie(expr: str) -> str:
    """缺数据时显示空白而不是错误值。"""
    return f'=IFERROR({expr},"")'


def build_template(example: dict | None = None) -> bytes:
    """example = {"name", "code", "annual"(DataFrame), "A"(Assumptions), "bs"(最新资产负债表 dict), "quote"}"""
    i18n.FORCE = "zh"          # 模板只有中文版
    try:
        return _build(example)
    finally:
        i18n.FORCE = None


def _build(example: dict | None) -> bytes:
    yr_cells = None
    NF = int(example["A"].years) if example and example.get("A") is not None else 5      # 预测期年数
    if example and example.get("drv") is not None and example.get("ovr"):
        # 网页上手动改过逐年假设：预测表中改为逐年数值（蓝色），否则仍用首末年插值公式
        d = example["drv"]
        yr_cells = {k: [(float(x), BLUE) for x in d[src]] for k, src in
                    (("g", "g"), ("m", "m"), ("other", "other"), ("da_pct", "da"), ("cx_pct", "capex"))}
    wb = Workbook()
    wb.remove(wb.active)
    an = example["annual"].tail(YEARS) if example else None
    H = f"'{HS}'!"

    # ── 使用说明
    ws = wb.create_sheet("使用说明")
    ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 96
    title = "折现法估值模板" + (f"（示例：{example['name']} {example['code']}）" if example else "")
    ws["A1"] = title; ws["A1"].font = TITLE
    steps = [
        ("第一步", "在「历史数据」表的蓝色格子里填入最近 5 个完整年度的报表科目（单位：万元），最早的年份放最左边。"
                 "科目名称与出处写在 G 列，可从 Wind、同花顺或年报直接摘取。"),
        ("第二步", "下方「计算指标」由公式自动算出，最后一行「偏差」应接近 0，用来检查数据是否录入正确。"),
        ("第三步", "在「假设」表的蓝色格子里填写预测假设。C 列「历史参考」由历史数据自动计算，可作为取值依据。"
                 "货币资金、交易性金融资产、有息负债建议用最新一期报表（可以是季报）。"),
        ("第四步", "「预测与折现法」「敏感性」全部由公式自动计算，无需填写。"),
        ("格式约定", "蓝色字体 = 输入；黑色字体 = 公式。只改蓝色格子。"),
        ("模型结构", "营业总收入 × 核心经营利润率 = 核心经营利润；+ 投资收益及其他 = 利润总额；"
                 "× (1 − 有效税率) × (1 − 少数股东占比) = 归母净利润。"
                 "折现法只对核心经营利润折现（投资收益默认不计入现金流），金融资产在净现金中按账面加回。"),
        ("自由现金流", "FCFF = 税后经营利润 + 折旧摊销 − 资本开支 − 营运资金增加；营运资金增加 = 营收增量 × 营运资本/营收。"),
        ("折现率", "股权成本 = 无风险利率 + 贝塔 × 市场风险溢价 + 规模溢价；WACC 按有息负债与总市值加权。"),
        ("折现时点", "第 1 年只计入估值基准日之后的部分（「假设」表的计入比例）；默认年中折现，现金流按各期中点折现。"),
        ("终值", "资本开支 / 营收从第 1 年线性过渡到稳定水平（预测末年，默认等于折旧摊销）；「终值年」= 预测末年按永续增长率"
               "外推一年，营运资金按永续增速计算；终值 = 终值年 FCFF ÷ (WACC − g)。预测期年数随网页设置（默认 5 年）。"),
        ("交叉检验", "「预测与折现法」表下方列出折现法隐含的 EV/EBITDA、市盈率与现价倍数，以及模型检查（通过 / 需关注 / 未通过）。"),
        ("单位", "金额：万元；股本：万股；每股：元；比率：%。"),
    ]
    for i, (k, v) in enumerate(steps, start=3):
        _c(ws, i, 1, k, BOLD, fill=KEY); _c(ws, i, 2, v, al="left")
        ws.row_dimensions[i].height = 32

    # ── 历史数据
    wh = wb.create_sheet(HS)
    wh.column_dimensions["A"].width = 26
    for c in COLS:
        wh.column_dimensions[c].width = 14
    wh.column_dimensions["G"].width = 70
    wh["A1"] = "历史数据（万元）"; wh["A1"].font = TITLE
    wh["A2"] = "只填蓝色格子；最早的年份放 B 列，最近的年份放 F 列。"; wh["A2"].font = NOTE
    _head(wh, 4, ["项目"] + [f"第{i + 1}年" for i in range(YEARS)] + ["报表与科目"])
    for r_, lab, src, key in INPUTS:
        _c(wh, r_, 1, lab, BOLD if r_ == 5 else BLACK, al="left")
        off = YEARS - len(an) if an is not None else 0        # 不足 5 年时靠右填，左侧留空
        for j, c in enumerate(COLS):
            v = None
            if an is not None and j >= off:
                row = an.iloc[j - off]
                v = int(an.index[j - off]) if r_ == 5 else (None if row.get(key) != row.get(key)
                                                              else float(row.get(key, 0.0)))
            _c(wh, r_, 2 + j, v, BLUE, "0" if r_ == 5 else N2, al="center" if r_ == 5 else None)
        _c(wh, r_, 7, src, NOTE, al="left")

    wh.cell(row=25, column=1, value="计算指标（公式，勿改）").font = BOLD
    calc = [
        (26, "营收增速（%）", lambda c, p: _ie(f"({c}6/{p}6-1)*100") if p else "", PCT),
        (27, "核心经营利润", lambda c, p: _ie(f"IF({c}6=\"\",NA(),{c}6-{c}7)"), N2),
        (28, "核心经营利润率（%）", lambda c, p: _ie(f"{c}27/{c}6*100"), PCT),
        (29, "投资收益及其他", lambda c, p: _ie(f"IF({c}8=\"\",NA(),{c}8-{c}27)"), N2),
        (30, "有效税率（%）", lambda c, p: _ie(f"{c}9/{c}8*100"), PCT),
        (31, "少数股东占比（%）", lambda c, p: _ie(f"(1-{c}11/{c}10)*100"), PCT),
        (32, "非经常性损益", lambda c, p: _ie(f"IF({c}12=\"\",NA(),{c}11-{c}12)"), N2),
        (33, "折旧摊销 / 营收（%）", lambda c, p: _ie(f"{c}13/{c}6*100"), PCT),
        (34, "资本开支 / 营收（%）", lambda c, p: _ie(f"{c}14/{c}6*100"), PCT),
        (35, "营运资本", lambda c, p: _ie(f"IF({c}6=\"\",NA(),{c}15+{c}16-{c}17)"), N2),
        (36, "营运资本 / 营收（%）", lambda c, p: _ie(f"{c}35/{c}6*100"), PCT),
        (37, "有息负债", lambda c, p: _ie(f"IF({c}6=\"\",NA(),{c}20+{c}21+{c}22+{c}23)"), N2),
        (38, "模型归母净利润", lambda c, p: _ie(f"({c}27+{c}29)*(1-{c}30/100)*(1-{c}31/100)"), N2),
        (39, "偏差（应接近 0）", lambda c, p: _ie(f"{c}38-{c}11"), N2),
    ]
    for r_, lab, fn, fm in calc:
        _c(wh, r_, 1, lab, BOLD if r_ == 39 else BLACK, al="left", fill=KEY if r_ == 39 else None)
        for j, c in enumerate(COLS):
            _c(wh, r_, 2 + j, fn(c, COLS[j - 1] if j else None), fmt=fm)
    note_rows = {35: "应收账款 + 存货 − 应付账款", 37: "短期借款 + 一年内到期的非流动负债 + 长期借款 + 应付债券",
                 39: "模型归母净利润 − 实际归母净利润，检验录入是否一致"}
    for r_, t in note_rows.items():
        _c(wh, r_, 7, t, NOTE, al="left")

    # ── 假设
    A = example["A"] if example else None
    bs = example["bs"] if example else {}
    q = example["quote"] if example else {}
    wa = wb.create_sheet(AS)
    wa.column_dimensions["A"].width = 34; wa.column_dimensions["B"].width = 14
    wa.column_dimensions["C"].width = 14; wa.column_dimensions["D"].width = 60
    wa["A1"] = "假设与输入"; wa["A1"].font = TITLE
    wa["A2"] = "只填 B 列蓝色格子；C 列由历史数据自动计算，供参考。"; wa["A2"].font = NOTE
    _head(wa, 3, ["项目", "假设值", "历史参考", "说明"])
    last3 = f"{COLS[-3]}{{r}}:{LAST}{{r}}"
    g = lambda k, d=None: (getattr(A, k) if A is not None else d)
    items = [
        # key, 名称, 假设值（None=公式）, 历史参考公式, 说明
        ("base_rev", "基期营业总收入（万元）", f"={H}{LAST}6", None, "取历史数据最近一年，无需填写"),
        ("g1", "第1年营收增速（%）", g("g_first"), _ie(f"{H}{LAST}26"), "参考：最近一年增速；中间年份线性插值"),
        ("g5", f"第{NF}年营收增速（%）", g("g_last"), _ie(f"(({H}{LAST}6/{H}B6)^(1/{YEARS - 1})-1)*100"),
         "参考：近 4 年复合增速；远期一般向行业长期增速收敛"),
        ("m1", "第1年核心经营利润率（%）", g("m_first"), _ie(f"{H}{LAST}28"), "参考：最近一年"),
        ("m5", f"第{NF}年核心经营利润率（%）", g("m_last"), _ie(f"AVERAGE({H}B28:{LAST}28)"), "参考：5 年平均"),
        ("o1", "第1年投资收益及其他（万元）", g("other_first"), _ie(f"{H}{LAST}29"),
         "利润总额 − 核心经营利润，含理财收益、政府补助等；参考：最近一年"),
        ("oc", "投资收益及其他年变化（%）", g("other_change", -15.0), None, "理财规模或收益率下行时取负值"),
        ("of", "其中计入经营现金流的比例（%）", g("other_in_fcf", 0.0), None,
         "0 = 折现法只对主业定价，金融资产在净现金中加回；若补助、退税属经常性可调高"),
        ("tax", "有效税率（%）", g("tax"), _ie(f"{H}{LAST}30"), "参考：最近一年"),
        ("mino", "少数股东占比（%）", g("minority"), _ie(f"{H}{LAST}31"), "参考：最近一年"),
        ("da", "折旧摊销 / 营收（%）", g("da_pct"), _ie(f"MEDIAN({H}{last3.format(r=33)})"), "参考：近 3 年中位数"),
        ("cx", "第1年资本开支 / 营收（%）", g("capex_pct"), _ie(f"MEDIAN({H}{last3.format(r=34)})"),
         f"参考：近 3 年中位数；中间年份线性过渡到第 {NF} 年"),
        ("nwc", "营运资本 / 营收（%）", g("nwc_pct"), _ie(f"MEDIAN({H}{last3.format(r=36)})"),
         "营运资金增加 = 营收增量 × 此比例；参考：近 3 年中位数"),
        ("rf", "无风险利率（%）", g("rf", 1.7), None, "10 年期国债收益率"),
        ("beta", "贝塔", g("beta", 1.0), None, "个股对沪深300 周收益率回归，或 Wind BETA 值"),
        ("erp", "市场风险溢价（%）", g("erp", 5.5), None, ""),
        ("sp", "规模溢价（%）", g("size_prem", 0.0), None, "小市值公司可加 0.5–1.0"),
        ("kd", "税前债务成本（%）", g("kd", 3.5), None, ""),
        ("gt", "永续增长率（%）", g("g_term", 2.5), None, "一般不高于长期名义 GDP 增速"),
        ("cx_ty", f"第{NF}年及终值年资本开支 / 营收（%）", g("capex_ty"), _ie(f"MEDIAN({H}{last3.format(r=33)})"),
         "稳定状态的维持性资本开支，一般取折旧摊销 / 营收；参考：折旧率近 3 年中位数"),
        ("mid", "年中折现（1=是，0=否）", int(A.mid_year) if A is not None else 1, None, "现金流按各期中点折现，估值报告通行做法"),
        ("stub", "第 1 年计入比例", g("stub", 1.0), None,
         "估值基准日之后的部分占第 1 个预测年的比例：基准日为基期年末时填 1，用半年报数据时填 0.5"),
        ("cash", "货币资金（万元）", bs.get("货币资金"), _ie(f"{H}{LAST}18"), "建议用最新一期报表；参考：最近年报"),
        ("fin", "交易性金融资产（万元）", bs.get("交易性金融资产"), _ie(f"{H}{LAST}19"), "同上"),
        ("debt", "有息负债（万元）", bs.get("有息负债"), _ie(f"{H}{LAST}37"), "同上"),
        ("add_cash", "货币资金计入净现金（1=是，0=否）", int(A.add_cash) if A is not None else 1, None, ""),
        ("add_fin", "交易性金融资产计入净现金（1=是，0=否）", int(A.add_fin) if A is not None else 1, None, ""),
        ("shares", "总股本（万股）", q["shares"] / 1e4 if q else None, None, ""),
        ("price", "现价（元）", q.get("price"), None, ""),
        ("mcap", "总市值（万元）", None, None, "总股本 × 现价，用于计算债务权重"),
    ]
    R = {}
    for i, (k, lab, v, ref, note) in enumerate(items, start=4):
        R[k] = f"'{AS}'!$B${i}"
        _c(wa, i, 1, lab, al="left")
        if k == "base_rev":
            _c(wa, i, 2, v, BLACK, N2)
        elif k == "mcap":
            _c(wa, i, 2, _ie(f"{R['shares']}*{R['price']}"), BLACK, N2)
        else:
            v = None if v is None else float(v)
            _c(wa, i, 2, v, BLUE, N2 if (v is not None and abs(v) >= 100) or k in ("cash", "fin", "debt", "shares", "o1")
               else "0.00")
        _c(wa, i, 3, ref, fmt=N2 if k in ("o1", "cash", "fin", "debt") else PCT)
        _c(wa, i, 4, note, NOTE, al="left")

    # ── 预测与折现法、敏感性（与网页导出共用同一套公式）
    base_year = int(an.index[-1]) if an is not None else None
    write_dcf(wb, R, FS, SS, f"{base_year}A" if base_year else "基期", {
        "g": (f"={H}{LAST}26", BLACK), "rev": (f"={R['base_rev']}", BLACK),
        "m": (f"={H}{LAST}28", BLACK), "core": (f"={H}{LAST}27", BLACK),
        "other": (f"={H}{LAST}29", BLACK), "pbt": (f"={H}{LAST}8", BLACK),
        "np": (f"={H}{LAST}11", BLACK), "da_pct": (f"={H}{LAST}33", BLACK),
        "cx_pct": (f"={H}{LAST}34", BLACK)}, base_year=base_year, n=NF, yr_cells=yr_cells,
        extra_checks=[("历史数据录入勾稽（各年偏差小于 1 万元）",
                       f'=IF(COUNT({H}B39:{LAST}39)=0,"需关注",IF(COUNTIF({H}B39:{LAST}39,">=1")'
                       f'+COUNTIF({H}B39:{LAST}39,"<=-1")=0,"通过","未通过"))',
                       "「历史数据」表最后一行；未填数据时显示需关注")])

    # 空白模板未填数据时，结果格显示为空而不是 #DIV/0!
    for w in (wb[FS], wb[SS]):
        for row in w.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("=") and not v.startswith("=IFERROR("):
                    cell.value = f'=IFERROR({v[1:]},"")'

    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
